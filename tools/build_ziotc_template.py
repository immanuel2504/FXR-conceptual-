from __future__ import annotations

from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT_PATH = Path(__file__).resolve().parents[1] / "Zebra IoT Connector Documentation Template.docx"


CHAPTERS: list[tuple[str, list[str]]] = [
    ("Overview", [
        "About Zebra IoT Connector",
        "How Zebra IoT Connector Works",
        "System Architecture",
        "System Components",
        "Interfaces",
        "Deployment Models",
        "Supported Readers and Versions",
        "Terminology",
    ]),
    ("Getting Started", [
        "Prerequisites",
        "Set Up the Reader",
        "Discover and Access the Reader",
        "Configure the Reader",
        "Verify Reader Status",
        "Quickstart Receive Your First Tag Data",
    ]),
    ("Reader Network Connectivity", [
        "Ethernet",
        "Wi Fi",
        "Hotspot and Access Point",
        "802 1X",
        "Bluetooth",
        "Cellular and eSIM",
        "Network Status",
    ]),
    ("Zebra IoT Connector Configuration", [
        "Configuration Workflow",
        "Endpoints and Interfaces",
        "Manage Endpoints",
        "Configure Interfaces",
        "Manage ZIoTC Service",
    ]),
    ("Backend and Cloud Connectivity", [
        "Generic MQTT Broker",
        "AWS IoT Core",
        "Azure IoT Hub",
        "HTTP and HTTPS",
        "Other Supported Destinations",
    ]),
    ("RFID Operations", [
        "Operating Modes",
        "Antenna and RF Configuration",
        "Tag Reporting",
        "Tag Filtering",
        "Start and Stop Conditions",
        "Tag Memory Operations",
        "Gen2X",
        "GPIO",
    ]),
    ("RFID Tag Data", [
        "Tag Data Structure",
        "Reader Identification",
        "MQTT Topics and Routing",
        "Receive Tag Data",
        "Filtering",
        "Batching",
        "Retention and Buffering",
        "Payload Examples",
    ]),
    ("Management Events", [
        "Heartbeat",
        "GPI and GPO Events",
        "Errors and Warnings",
        "Notifications",
        "Firmware Update Progress",
        "User Application Events",
        "Payload Examples",
    ]),
    ("BLE Scanning and Events", [
        "Configure BLE Scanning",
        "Receive BLE Events",
        "BLE Event Structure",
        "BLE Data Routing",
        "Payload Examples",
    ]),
    ("Application Integration", [
        "Application to Reader Architecture",
        "Choose MQTT or REST",
        "MQTT Application Integration",
        "REST Application Integration",
        "Process Reader Data",
        "Sample Applications",
    ]),
    ("Deployment and Scaling", [
        "Single Reader Deployment",
        "Multiple Reader Architecture",
        "Shared MQTT Broker",
        "Centralized Backend Architecture",
        "Reader Identification",
        "Provisioning",
        "Scaling Considerations",
    ]),
    ("Security", [
        "Authentication",
        "Network Security",
        "MQTT Security",
        "REST Security",
        "Certificates",
        "TLS and mTLS",
        "Security Best Practices",
    ]),
    ("Monitoring and Maintenance", [
        "Reader Health",
        "ZIoTC Status",
        "Endpoint and Interface Status",
        "Logs and Diagnostics",
        "Backup and Restore",
        "Firmware Updates",
    ]),
    ("Guides and Best Practices", [
        "Configure a Reader for MQTT",
        "Send Tag Data to an MQTT Broker",
        "Connect to AWS IoT Core",
        "Connect to Azure IoT Hub",
        "Configure RFID Operations",
        "Receive and Process Tag Data",
        "Send Commands and Process Responses",
        "Configure Management Events",
        "Configure BLE Data Routing",
        "Integrate Multiple Readers",
        "Reader Configuration Best Practices",
        "MQTT Topic Design Best Practices",
        "Security Best Practices",
        "Deployment Best Practices",
    ]),
    ("Troubleshooting", [
        "Reader Access Issues",
        "Network Issues",
        "ZIoTC Configuration Issues",
        "MQTT Connectivity Issues",
        "Cloud Connectivity Issues",
        "Authentication and Certificate Issues",
        "RFID Issues",
        "BLE Issues",
        "Missing Tag Data",
        "Missing Events",
        "Command and Response Issues",
    ]),
    ("Solution Examples and Use Cases", [
        "Warehouse Inventory Tracking",
        "Dock Door Shipment Verification",
        "Asset Tracking",
        "Retail Inventory",
        "Reader to MQTT Broker to WMS",
        "Reader to AWS IoT Core",
        "Reader to Azure IoT Hub",
        "Multiple Readers to Central Application",
    ]),
    ("Migration", [
        "Migration Overview",
        "FxConnect to ZIoTC",
        "123RFID Desktop to ZIoTC",
        "RFID API3 to ZIoTC",
        "Configuration Mapping",
    ]),
    ("API Reference", [
        "REST API Reference",
        "MQTT API Reference",
        "MQTT Command Reference",
        "MQTT Topic Reference",
        "Tag Data Schema",
        "Management Event Schema",
        "Heartbeat Schema",
        "BLE Event Schema",
        "Configuration Parameters",
        "Defaults and Limits",
        "Supported Features by Reader",
        "Compatibility Matrix",
        "Status and Error Codes",
        "Known Issues",
        "Release Notes",
    ]),
]


def set_font(run, name: str, size: float | None = None, *, bold: bool | None = None,
             italic: bool | None = None, color: str | None = None) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def set_style_font(style, name: str, size: float, *, bold: bool = False,
                   italic: bool = False, color: str = "000000") -> None:
    style.font.name = name
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.italic = italic
    style.font.color.rgb = RGBColor.from_string(color)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), name)
    rfonts.set(qn("w:hAnsi"), name)
    rfonts.set(qn("w:eastAsia"), name)
    color_el = rpr.find(qn("w:color"))
    if color_el is None:
        color_el = OxmlElement("w:color")
        rpr.append(color_el)
    color_el.set(qn("w:val"), color)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def set_cell_margins(cell, top: int = 90, start: int = 110,
                     bottom: int = 90, end: int = 110) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        element = tc_mar.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            tc_mar.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def set_table_borders(table, color: str = "D9D9D9", size: int = 6) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), str(size))
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), color)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def add_field(paragraph, instruction: str, result: str = "") -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = result
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])


def enable_field_updates(doc: Document) -> None:
    settings = doc.settings._element
    existing = settings.find(qn("w:updateFields"))
    if existing is None:
        existing = OxmlElement("w:updateFields")
        settings.append(existing)
    existing.set(qn("w:val"), "true")


def configure_heading_numbering(doc: Document) -> int:
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(x.get(qn("w:abstractNumId"))) for x in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=0) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "multilevel")
    abstract.append(multi)

    levels = [
        (0, "Heading1", "%1", 0, 0),
        (1, "Heading2", "%1.%2", 360, 360),
    ]
    for ilvl, style_id, lvl_text, left, hanging in levels:
        lvl = OxmlElement("w:lvl")
        lvl.set(qn("w:ilvl"), str(ilvl))
        start = OxmlElement("w:start")
        start.set(qn("w:val"), "1")
        num_fmt = OxmlElement("w:numFmt")
        num_fmt.set(qn("w:val"), "decimal")
        p_style = OxmlElement("w:pStyle")
        p_style.set(qn("w:val"), style_id)
        text = OxmlElement("w:lvlText")
        text.set(qn("w:val"), lvl_text)
        suff = OxmlElement("w:suff")
        suff.set(qn("w:val"), "space")
        lvl.append(start)
        lvl.append(num_fmt)
        lvl.append(p_style)
        lvl.append(text)
        lvl.append(suff)
        p_pr = OxmlElement("w:pPr")
        ind = OxmlElement("w:ind")
        ind.set(qn("w:left"), str(left))
        ind.set(qn("w:hanging"), str(hanging))
        p_pr.append(ind)
        lvl.append(p_pr)
        abstract.append(lvl)
    numbering.insert(0, abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_num_id = OxmlElement("w:abstractNumId")
    abstract_num_id.set(qn("w:val"), str(abstract_id))
    num.append(abstract_num_id)
    numbering.append(num)
    return num_id


def apply_numbering_to_style(style, num_id: int, ilvl: int) -> None:
    p_pr = style.element.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)
    ilvl_el = num_pr.find(qn("w:ilvl"))
    if ilvl_el is None:
        ilvl_el = OxmlElement("w:ilvl")
        num_pr.append(ilvl_el)
    ilvl_el.set(qn("w:val"), str(ilvl))
    num_id_el = num_pr.find(qn("w:numId"))
    if num_id_el is None:
        num_id_el = OxmlElement("w:numId")
        num_pr.append(num_id_el)
    num_id_el.set(qn("w:val"), str(num_id))


def set_keep_with_next(paragraph) -> None:
    paragraph.paragraph_format.keep_with_next = True


def add_header_footer(section) -> None:
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)

    header_p = section.header.paragraphs[0]
    header_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header_p.paragraph_format.space_after = Pt(0)
    header_p.paragraph_format.tab_stops.add_tab_stop(Inches(6.7))
    left = header_p.add_run("Zebra IoT Connector")
    set_font(left, "Arial", 8.5, bold=True, color="333333")
    right = header_p.add_run("\tDocumentation Template")
    set_font(right, "Arial", 8.5, color="666666")

    footer_p = section.footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_p.paragraph_format.space_before = Pt(0)
    run = footer_p.add_run("Page ")
    set_font(run, "Arial", 8.5, color="666666")
    add_field(footer_p, " PAGE ", "1")
    run = footer_p.add_run(" of ")
    set_font(run, "Arial", 8.5, color="666666")
    add_field(footer_p, " NUMPAGES ", "1")
    for r in footer_p.runs:
        set_font(r, "Arial", 8.5, color="666666")


def add_page_number_restart(section, start: int = 1) -> None:
    sect_pr = section._sectPr
    pg_num = sect_pr.find(qn("w:pgNumType"))
    if pg_num is None:
        pg_num = OxmlElement("w:pgNumType")
        sect_pr.append(pg_num)
    pg_num.set(qn("w:start"), str(start))


def set_section_page(section) -> None:
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.9)
    section.right_margin = Inches(0.9)
    section.gutter = Inches(0)


def add_front_heading(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="Front Matter Heading")
    p.add_run(text)


def add_bullet(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.first_line_indent = Inches(-0.15)
    p.paragraph_format.space_after = Pt(4)
    p.add_run(text)


def authoring_prompt(chapter: str, section: str) -> str:
    c = chapter.lower()
    s = section.lower()
    if "troubleshooting" in c:
        return ("[Authoring prompt] Describe the symptom, likely causes, diagnostic checks, "
                "corrective actions, expected recovery, and escalation information.")
    if "api reference" in c or any(word in s for word in ("schema", "reference", "parameters", "codes", "matrix", "limits")):
        return ("[Authoring prompt] Document identifiers, fields or parameters, data types, required values, "
                "defaults, limits, responses, errors, compatibility, and a validated example as applicable.")
    if "guides" in c or any(word in s for word in ("configure", "send ", "connect ", "receive ", "integrate ")):
        return ("[Authoring prompt] State the goal and prerequisites, list the required values, provide numbered "
                "steps, show the expected result, and explain how to verify or reverse the change.")
    if "security" in c or any(word in s for word in ("authentication", "certificate", "tls")):
        return ("[Authoring prompt] Explain the control, trust boundary, credentials or certificates, supported "
                "options, configuration steps, validation method, and security implications.")
    if "solution examples" in c:
        return ("[Authoring prompt] Describe the scenario, architecture, components, data flow, configuration "
                "decisions, sample messages, validation steps, and operational considerations.")
    if "migration" in c:
        return ("[Authoring prompt] Define the source and target, prerequisites, feature mapping, migration steps, "
                "validation criteria, limitations, and rollback guidance.")
    if "deployment" in c:
        return ("[Authoring prompt] Describe the topology, components, identity model, provisioning approach, "
                "capacity assumptions, failure handling, and operational tradeoffs.")
    if "monitoring" in c:
        return ("[Authoring prompt] Define the health signal, where to view or collect it, normal and abnormal "
                "states, diagnostic steps, maintenance actions, and verification criteria.")
    if any(word in c for word in ("tag data", "events")) or "payload" in s or "routing" in s:
        return ("[Authoring prompt] Explain the source, transport or topic, message structure, required fields, "
                "routing behavior, buffering rules, filtering behavior, and a representative payload.")
    if "rfid operations" in c:
        return ("[Authoring prompt] Explain the feature, supported settings, defaults and limits, configuration "
                "procedure, operational effect, validation method, and a practical example.")
    if "connectivity" in c or "network" in s or any(word in s for word in ("mqtt", "http", "aws", "azure")):
        return ("[Authoring prompt] Document prerequisites, endpoints, ports and protocols, authentication, "
                "certificates, configuration steps, status indicators, and connectivity tests.")
    if "overview" in c:
        return ("[Authoring prompt] Define the concept, explain why it matters, show its relationship to other "
                "components, state supported scope, and link to detailed procedures or references.")
    return ("[Authoring prompt] Explain the purpose and scope, prerequisites, configuration or behavior, "
            "defaults and limitations, verification method, and related references.")


def add_document_control_table(doc: Document) -> None:
    rows = [
        ("Product or firmware version", "[Enter supported version or range]"),
        ("Document version", "[Enter document version]"),
        ("Audience", "[Enter primary audience and assumed experience]"),
        ("Document owner", "[Enter owner or team]"),
        ("Status", "[Draft | In review | Approved]"),
        ("Last updated", "[DD Month YYYY]"),
    ]
    table = doc.add_table(rows=len(rows), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(2.2)
    table.columns[1].width = Inches(4.25)
    set_table_borders(table)
    for i, (label, value) in enumerate(rows):
        for cell in table.rows[i].cells:
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(table.cell(i, 0), "F2F2F2")
        p = table.cell(i, 0).paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(label)
        set_font(r, "Arial", 9.5, bold=True)
        p = table.cell(i, 1).paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(value)
        set_font(r, "Arial", 9.5, color="666666", italic=True)


def add_revision_table(doc: Document) -> None:
    table = doc.add_table(rows=2, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [Inches(1.0), Inches(1.3), Inches(1.55), Inches(3.0)]
    headers = ["Version", "Date", "Author", "Summary of changes"]
    values = ["[0.1]", "[DD Mon YYYY]", "[Name]", "[Describe the change and approval status]"]
    set_table_borders(table)
    set_repeat_table_header(table.rows[0])
    for col, width in enumerate(widths):
        for row in table.rows:
            row.cells[col].width = width
    for j, header in enumerate(headers):
        cell = table.cell(0, j)
        set_cell_shading(cell, "333333")
        set_cell_margins(cell, top=105, bottom=105)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(header)
        set_font(r, "Arial", 9, bold=True, color="FFFFFF")
    for j, value in enumerate(values):
        cell = table.cell(1, j)
        set_cell_margins(cell, top=105, bottom=105)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if j < 3 else WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(value)
        set_font(r, "Arial", 9, italic=True, color="666666")


def add_topic_pattern_table(doc: Document) -> None:
    data = [
        ("Purpose and scope", "Open with what the topic covers, who needs it, and the outcome."),
        ("Prerequisites", "List required hardware, software, permissions, network access, and source values."),
        ("Concept or architecture", "Explain relationships and data flow before detailed configuration."),
        ("Procedure", "Use numbered steps with exact fields, commands, values, and expected results."),
        ("Verification", "State how the reader confirms that the task or configuration succeeded."),
        ("Examples", "Use realistic topics, payloads, requests, responses, and diagrams where they help."),
        ("Troubleshooting", "Connect symptoms to diagnostic evidence, corrective action, and escalation."),
    ]
    table = doc.add_table(rows=1 + len(data), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(1.75)
    table.columns[1].width = Inches(4.75)
    set_table_borders(table)
    set_repeat_table_header(table.rows[0])
    for j, header in enumerate(("Topic element", "Authoring guidance")):
        cell = table.cell(0, j)
        set_cell_shading(cell, "333333")
        set_cell_margins(cell, top=105, bottom=105)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(header)
        set_font(r, "Arial", 9.5, bold=True, color="FFFFFF")
    for i, (label, guidance) in enumerate(data, start=1):
        if i % 2 == 0:
            for cell in table.rows[i].cells:
                set_cell_shading(cell, "F7F7F7")
        for cell in table.rows[i].cells:
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = table.cell(i, 0).paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(label)
        set_font(r, "Arial", 9, bold=True)
        p = table.cell(i, 1).paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(guidance)
        set_font(r, "Arial", 9)


def build_document() -> None:
    doc = Document()
    for section in doc.sections:
        set_section_page(section)

    props = doc.core_properties
    props.title = "Zebra IoT Connector Documentation Template"
    props.subject = "Reusable structure for Zebra IoT Connector product documentation"
    props.author = ""
    props.last_modified_by = ""
    props.keywords = "Zebra IoT Connector, ZIoTC, RFID, MQTT, REST, BLE, documentation"
    props.comments = "Reusable documentation template"

    styles = doc.styles
    normal = styles["Normal"]
    set_style_font(normal, "Arial", 10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12

    title = styles["Title"]
    set_style_font(title, "Arial", 28, bold=True)
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(14)
    title.paragraph_format.keep_with_next = True
    title_p_pr = title.element.get_or_add_pPr()
    title_border = title_p_pr.find(qn("w:pBdr"))
    if title_border is not None:
        title_p_pr.remove(title_border)

    subtitle = styles.add_style("Document Subtitle", WD_STYLE_TYPE.PARAGRAPH)
    set_style_font(subtitle, "Arial", 13, color="000000")
    subtitle.paragraph_format.space_after = Pt(20)

    front = styles.add_style("Front Matter Heading", WD_STYLE_TYPE.PARAGRAPH)
    set_style_font(front, "Arial", 19, bold=True)
    front.paragraph_format.space_before = Pt(0)
    front.paragraph_format.space_after = Pt(12)
    front.paragraph_format.keep_with_next = True

    prompt = styles.add_style("Authoring Prompt", WD_STYLE_TYPE.PARAGRAPH)
    set_style_font(prompt, "Arial", 9.25, italic=True, color="666666")
    prompt.paragraph_format.left_indent = Inches(0.18)
    prompt.paragraph_format.space_after = Pt(8)
    prompt.paragraph_format.line_spacing = 1.05

    small = styles.add_style("Small Instruction", WD_STYLE_TYPE.PARAGRAPH)
    set_style_font(small, "Arial", 8.5, italic=True, color="666666")
    small.paragraph_format.space_before = Pt(6)
    small.paragraph_format.space_after = Pt(6)

    h1 = styles["Heading 1"]
    set_style_font(h1, "Arial", 19, bold=True)
    h1.paragraph_format.space_before = Pt(0)
    h1.paragraph_format.space_after = Pt(8)
    h1.paragraph_format.keep_with_next = True
    h1.paragraph_format.page_break_before = True

    h2 = styles["Heading 2"]
    set_style_font(h2, "Arial", 12.5, bold=True)
    h2.paragraph_format.space_before = Pt(9)
    h2.paragraph_format.space_after = Pt(3)
    h2.paragraph_format.keep_with_next = True

    num_id = configure_heading_numbering(doc)
    apply_numbering_to_style(h1, num_id, 0)
    apply_numbering_to_style(h2, num_id, 1)

    enable_field_updates(doc)

    # Cover page
    cover = doc.sections[0]
    set_section_page(cover)
    cover.different_first_page_header_footer = True
    for _ in range(4):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
    p = doc.add_paragraph(style="Small Instruction")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run("DOCUMENTATION TEMPLATE")
    set_font(r, "Arial", 9.5, bold=True, color="333333")
    p = doc.add_paragraph(style="Title")
    p.add_run("Zebra IoT Connector")
    p = doc.add_paragraph(style="Document Subtitle")
    p.add_run("Technical documentation structure and authoring framework")
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(10)
    p.add_run(
        "Use this template to create a consistent product guide for setup, connectivity, RFID operations, "
        "data integration, security, maintenance, troubleshooting, migration, and API reference content."
    )
    p = doc.add_paragraph(style="Small Instruction")
    p.paragraph_format.space_before = Pt(80)
    p.add_run("Replace every bracketed instruction before publication and refresh all fields in Microsoft Word.")

    # New main section after the cover.
    main = doc.add_section(WD_SECTION.NEW_PAGE)
    set_section_page(main)
    main.different_first_page_header_footer = False
    main.header.is_linked_to_previous = False
    main.footer.is_linked_to_previous = False
    add_page_number_restart(main, 1)

    add_front_heading(doc, "Document Control")
    p = doc.add_paragraph()
    p.add_run(
        "Complete these fields before content review so contributors understand the supported product scope, "
        "intended audience, ownership, and approval state."
    )
    add_document_control_table(doc)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    r = p.add_run("Source of truth: ")
    set_font(r, "Arial", 10, bold=True)
    r = p.add_run("[Enter the authoritative repository, publication system, or owner]")
    set_font(r, "Arial", 10, italic=True, color="666666")

    doc.add_page_break()
    add_front_heading(doc, "Revision History")
    p = doc.add_paragraph()
    p.add_run("Record material content and structure changes that reviewers and publishers need to trace.")
    add_revision_table(doc)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16)
    r = p.add_run("Approval record: ")
    set_font(r, "Arial", 10, bold=True)
    r = p.add_run("[Enter approver, approval date, and release or publication identifier]")
    set_font(r, "Arial", 10, italic=True, color="666666")

    doc.add_page_break()
    add_front_heading(doc, "How to Use This Template")
    p = doc.add_paragraph()
    p.add_run(
        "The chapter and topic headings already use Word heading styles and automatic numbering. Keep those "
        "styles so navigation, numbering, and the table of contents remain reliable."
    )
    add_bullet(doc, "Replace each bracketed authoring prompt with reader-ready content or remove the topic if it is out of scope.")
    add_bullet(doc, "Use short introductions before procedures, tables, code samples, diagrams, or payload examples.")
    add_bullet(doc, "State product, firmware, protocol, and cloud-service versions whenever behavior depends on them.")
    add_bullet(doc, "Use numbered steps for procedures and include an observable result or verification step.")
    add_bullet(doc, "Before publication, press Ctrl+A and then F9 in Microsoft Word to refresh numbering, page references, and the table of contents.")
    add_topic_pattern_table(doc)

    doc.add_page_break()
    add_front_heading(doc, "Contents")
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(10)
    add_field(p, ' TOC \\o "1-2" \\h \\z \\u ', "Right click and select Update Field to build the table of contents")
    p = doc.add_paragraph(style="Small Instruction")
    p.add_run("Refresh this table after adding, renaming, or removing topics: press Ctrl+A, then F9 in Microsoft Word.")

    # Main manual structure.
    for chapter, sections in CHAPTERS:
        p = doc.add_paragraph(style="Heading 1")
        p.add_run(chapter)
        intro = doc.add_paragraph(style="Authoring Prompt")
        intro.add_run(
            "[Chapter overview] Summarize what readers will learn, who should read this chapter, "
            "and any prerequisites or related chapters."
        )
        for section_name in sections:
            p = doc.add_paragraph(style="Heading 2")
            p.add_run(section_name)
            p = doc.add_paragraph(style="Authoring Prompt")
            p.add_run(authoring_prompt(chapter, section_name))

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT_PATH)
    print(OUTPUT_PATH)


if __name__ == "__main__":
    build_document()
