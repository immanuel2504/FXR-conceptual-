from pathlib import Path
from zipfile import ZipFile

from docx import Document


docx_path = Path(__file__).resolve().parents[1] / "Zebra IoT Connector Documentation Template.docx"
document = Document(docx_path)

heading_1_count = sum(p.style.name == "Heading 1" for p in document.paragraphs)
heading_2_count = sum(p.style.name == "Heading 2" for p in document.paragraphs)

with ZipFile(docx_path) as package:
    names = package.namelist()
    document_xml = package.read("word/document.xml").decode("utf-8")
    settings_xml = package.read("word/settings.xml").decode("utf-8")
    styles_xml = package.read("word/styles.xml").decode("utf-8")

assert heading_1_count == 19, heading_1_count
assert heading_2_count == 148, heading_2_count
assert len(document.tables) == 3, len(document.tables)
assert "TOC \\o" in document_xml and "\\h \\z \\u" in document_xml
assert "word/numbering.xml" in names

title_style = styles_xml.split(
    '<w:style w:type="paragraph" w:styleId="Title"', 1
)[1].split("</w:style>", 1)[0]
assert "<w:pBdr>" not in title_style

print(
    {
        "file": str(docx_path),
        "bytes": docx_path.stat().st_size,
        "paragraphs": len(document.paragraphs),
        "tables": len(document.tables),
        "heading_1": heading_1_count,
        "heading_2": heading_2_count,
        "sections": len(document.sections),
        "zip_entries": len(names),
        "toc": True,
        "update_fields": "updateFields" in settings_xml,
        "title_border": False,
    }
)
